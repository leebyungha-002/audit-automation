"""사람 검토 단계. 검토용 엑셀을 내보내고, 사용자가 수정한 파일을 다시 읽어 DB에 반영한다.

- 수정 가능한 열은 머리글이 노란색이다. finding_id와 회색 열은 고치지 않는다.
- AI 원본(llm_original_json)은 재반영 후에도 그대로 보존된다.
- 원문 발췌나 기준서를 고치면 원문과 다시 대조해 검증 상태를 갱신한다.
"""
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .db import Store, now
from .models import FRAMEWORKS, TARGETS
from .structure import VERIFIED, verify_excerpt, verify_standards

SHEET = "지적사항 검토"
STATUSES = ["미검토", "확정", "제외"]

# (머리글, DB 필드, 열 너비, 수정 가능 여부)
COLUMNS = [
    ("finding_id", "finding_id", 16, False),
    ("검토상태", "review_status", 10, True),
    ("검토자 메모", "reviewer_note", 24, True),
    ("제목", "title", 28, True),
    ("발행기관", "issuer", 9, False),
    ("결정연도", "year", 9, True),
    ("쟁점분야", "issue_area", 16, True),
    ("지적대상", "finding_target", 13, True),
    ("지적유형", "finding_type", 15, True),
    ("관련 기준서", "standards", 34, True),
    ("관련 계정", "related_accounts", 24, True),
    ("위험 요약", "risk_summary", 60, True),
    ("원문 발췌", "source_excerpt", 60, True),
    ("발췌 검증", "excerpt_status", 12, False),
    ("권장 확인사항 (AI 제안, 원문 아님)", "audit_hint", 60, True),
    ("기준서 검증", "standards_check", 26, False),
    ("회계결산일", "fiscal_period", 20, False),
    ("출처 파일", "source_file", 34, False),
    ("페이지", "source_page", 8, False),
]
EDITABLE = {field for _, field, _, editable in COLUMNS if editable}
STANDARD_LINE = re.compile(rf"^({'|'.join(map(re.escape, FRAMEWORKS))})\s+(.+?)(?:\((.*)\))?$")

GUIDE = [
    "감리 지적사항 검토용 파일",
    "",
    "1. '지적사항 검토' 시트에서 머리글이 노란색인 열만 수정합니다. finding_id와 회색 열은 고치지 않습니다.",
    "2. 검토상태: 미검토 / 확정 / 제외 중에서 고릅니다. '제외'는 체크리스트에 나오지 않습니다.",
    "3. 관련 기준서: 한 줄에 하나씩 '구분 번호(명칭)' 형식. 구분은 K-IFRS / 일반기업회계기준 / 회계감사기준 / 기타.",
    "   예) K-IFRS 제1115호(고객과의 계약에서 생기는 수익)",
    "4. 관련 계정: 쉼표로 구분합니다. 예) 매출, 매출채권",
    "5. 원문 발췌를 고치면 재반영할 때 원문과 다시 대조합니다. '발췌 검증'이 불일치인 항목은 원문을 직접 확인하세요.",
    "6. '권장 확인사항'은 AI가 제안한 것으로 원문에 없는 내용입니다. 반드시 감사인이 검토·수정합니다.",
    "7. 저장 후 'python main.py review-import' 를 실행하면 DB에 반영됩니다. AI가 처음 만든 원본은 DB에 따로 보존됩니다.",
]


def _standards_to_text(standards: list[dict]) -> str:
    return "\n".join(f"{s['framework']} {s['number']}" + (f"({s['name']})" if s.get("name") else "") for s in standards)


def _standards_check_to_text(checked: list[dict]) -> str:
    return "\n".join(f"{s['number']}: {s['check']}" for s in checked)


def _cell_value(row, field: str):
    value = row[field]
    if field == "standards":
        return _standards_to_text(json.loads(value or "[]"))
    if field == "standards_check":
        return _standards_check_to_text(json.loads(value or "[]"))
    if field == "related_accounts":
        return ", ".join(json.loads(value or "[]"))
    return value


def export_review(cfg: dict, store: Store, log, status: str | None = None) -> Path | None:
    sql = "SELECT * FROM findings"
    params: tuple = ()
    if status:
        sql, params = sql + " WHERE review_status=?", (status,)
    rows = store.query(sql + " ORDER BY issuer, finding_id", params)
    if not rows:
        log.warning("내보낼 지적사항이 없습니다.")
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_edit = PatternFill("solid", fgColor="FFF2CC")
    fill_lock = PatternFill("solid", fgColor="D9D9D9")
    fill_flag = PatternFill("solid", fgColor="F8CBAD")

    for col, (header, _, width, editable) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = fill_edit if editable else fill_lock
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32

    for r, row in enumerate(rows, start=2):
        for col, (_, field, _, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=col, value=_cell_value(row, field))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if field == "excerpt_status" and row["excerpt_status"] not in VERIFIED:
                cell.fill = fill_flag
            if field == "standards_check" and "원문 미확인" in (cell.value or ""):
                cell.fill = fill_flag

    last = len(rows) + 1
    for header, options in (("검토상태", STATUSES), ("지적대상", TARGETS), ("지적유형", cfg["finding_types"])):
        letter = get_column_letter(next(i for i, c in enumerate(COLUMNS, start=1) if c[0] == header))
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "목록에 있는 값만 입력할 수 있습니다.", "입력 오류"
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"

    guide = wb.create_sheet("안내")
    for i, line in enumerate(GUIDE, start=1):
        guide.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    guide.column_dimensions["A"].width = 120

    path = cfg["paths"]["output"] / f"review_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    log.info("검토용 엑셀 내보냄: %s (%d건)", path, len(rows))
    store.log("review", "INFO", path.name, f"검토용 엑셀 내보냄 {len(rows)}건")
    return path


def _norm(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("_x000D_", "").strip()


def _parse_standards(text: str) -> list[dict] | None:
    """'구분 번호(명칭)' 줄들을 목록으로 바꾼다. 형식이 틀린 줄이 있으면 None."""
    items = []
    for line in filter(None, (ln.strip() for ln in text.split("\n"))):
        m = STANDARD_LINE.match(line)
        if not m:
            return None
        items.append({"framework": m.group(1), "number": m.group(2).strip(), "name": (m.group(3) or "").strip() or None})
    return items


def import_review(cfg: dict, store: Store, log, path: Path | None = None) -> None:
    if path is None:
        candidates = sorted(cfg["paths"]["output"].glob("review_*.xlsx"))
        if not candidates:
            log.error("output 폴더에 review_*.xlsx 파일이 없습니다.")
            return
        path = candidates[-1]
    try:
        ws = load_workbook(path, data_only=True)[SHEET]
    except (KeyError, OSError) as e:
        log.error("검토 파일을 읽을 수 없습니다: %s (%s)", path, e)
        return

    header_to_field = {h: f for h, f, _, _ in COLUMNS}
    col_field = {i: header_to_field[c.value] for i, c in enumerate(ws[1]) if c.value in header_to_field}
    if "finding_id" not in col_field.values():
        log.error("finding_id 열이 없습니다. 내보낸 파일의 머리글을 바꾸지 마세요.")
        return

    n_changed = n_fields = 0
    status_count = {s: 0 for s in STATUSES}
    for values in ws.iter_rows(min_row=2, values_only=True):
        data = {col_field[i]: v for i, v in enumerate(values) if i in col_field}
        fid = _norm(data.get("finding_id"))
        if not fid:
            continue
        found = store.query("SELECT * FROM findings WHERE finding_id=?", (fid,))
        if not found:
            log.warning("%s: DB에 없는 finding_id — 건너뜀", fid)
            store.log("review", "WARNING", path.name, f"{fid}: DB에 없는 finding_id")
            continue
        old = found[0]
        updates: dict = {}

        for field in EDITABLE & data.keys():
            new_text, old_text = _norm(data[field]), _norm(_cell_value(old, field))
            if new_text == old_text:
                continue
            if field == "review_status":
                if new_text not in STATUSES:
                    log.warning("%s: 검토상태 '%s'는 허용값이 아님 — 무시", fid, new_text)
                    continue
                updates[field] = new_text
            elif field == "finding_target":
                if new_text and new_text not in TARGETS:
                    log.warning("%s: 지적대상 '%s'는 허용값이 아님 — 무시", fid, new_text)
                    continue
                updates[field] = new_text or None
            elif field == "year":
                if new_text and not new_text.isdigit():
                    log.warning("%s: 결정연도 '%s'는 숫자가 아님 — 무시", fid, new_text)
                    continue
                updates[field] = int(new_text) if new_text else None
            elif field == "standards":
                parsed = _parse_standards(new_text)
                if parsed is None:
                    log.warning("%s: 관련 기준서 형식 오류 — 무시 (한 줄에 '구분 번호(명칭)')", fid)
                    store.log("review", "WARNING", path.name, f"{fid}: 관련 기준서 형식 오류로 미반영")
                    continue
                updates[field] = parsed
            elif field == "related_accounts":
                updates[field] = [a.strip() for a in re.split(r"[,\n]", new_text) if a.strip()]
            else:
                updates[field] = new_text or None

        if "source_excerpt" in updates or "standards" in updates:
            seg = store.query("SELECT pages_json FROM segments WHERE segment_id=?", (old["segment_id"],))
            pages = json.loads(seg[0]["pages_json"]) if seg else []
            if "source_excerpt" in updates and pages:
                status, ratio, page = verify_excerpt(updates["source_excerpt"], pages,
                                                     cfg["verify"]["fuzzy_threshold"], old["extract_mode"])
                updates["excerpt_status"], updates["excerpt_match_ratio"] = status, ratio
                if page:
                    updates["source_page"] = page
                if status not in VERIFIED:
                    log.warning("%s: 수정한 원문 발췌가 원문과 일치하지 않음 (%s, %s)", fid, status, ratio)
                    store.log("review", "WARNING", path.name, f"{fid}: 수정한 발췌 검증 {status}({ratio})")
            if "standards" in updates and pages:
                updates["standards_check"] = verify_standards(updates["standards"], pages, old["extract_mode"])

        final_status = updates.get("review_status", old["review_status"])
        status_count[final_status] = status_count.get(final_status, 0) + 1
        final_excerpt_status = updates.get("excerpt_status", old["excerpt_status"])
        if final_status == "확정" and final_excerpt_status not in VERIFIED:
            log.warning("%s: 확정 처리되었으나 발췌 검증이 '%s' — 체크리스트 출력 전 원문 확인 필요", fid, final_excerpt_status)
        if not updates:
            continue

        updates["reviewed_at"] = now()
        sets = ", ".join(f"{k}=?" for k in updates)
        params = [json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v for v in updates.values()]
        store.conn.execute(f"UPDATE findings SET {sets} WHERE finding_id=?", (*params, fid))
        changed = [k for k in updates if k in EDITABLE]
        store.log("review", "INFO", path.name, f"{fid}: 수정 반영 ({', '.join(changed)})")
        log.info("%s: 반영 (%s)", fid, ", ".join(changed))
        n_changed += 1
        n_fields += len(changed)

    store.commit()
    log.info("재반영 완료: %s | 변경 %d건(항목 %d개) | 미검토 %d, 확정 %d, 제외 %d",
             path.name, n_changed, n_fields, status_count["미검토"], status_count["확정"], status_count["제외"])
