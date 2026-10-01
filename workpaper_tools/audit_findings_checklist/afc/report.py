"""체크리스트 엑셀 출력. 표준 계정 분류별로 회사 계정과 감리 지적사항을 묶어 보여준다.

환각 방지 규칙: 원문 발췌가 없거나 원문 대조를 통과하지 못한 지적사항은 체크리스트에 싣지 않고
처리 로그 시트에 사유를 남긴다. 모든 위험 항목에는 finding_id·출처 파일·페이지·발췌가 붙는다.
"""
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .db import Store
from .signals import load_signals
from .structure import VERIFIED

AUDITOR_STATUSES = ["미검토", "검토중", "검토완료", "해당없음"]
NOTICE = [
    "※ 본 체크리스트는 금융감독원·한국공인회계사회가 공개한 감리 지적사례를 회사 계정과 대사한 참고자료입니다.",
    "※ 우선순위는 지적 빈도와 최근 연도를 기준으로 한 힌트일 뿐이며, 최종 위험평가는 감사인의 판단에 따릅니다.",
    "※ '권장 확인사항'은 AI가 제안한 것으로 원문에 없는 내용입니다. '주요 위험'과 '근거'는 원문에 기반하며 finding_id로 추적할 수 있습니다.",
]
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
CELL_LIMIT = 32000  # 엑셀 셀 글자 수 한도(32,767) 아래로 유지

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_HEAD = PatternFill("solid", fgColor="DDEBF7")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_NONE = PatternFill("solid", fgColor="F2F2F2")


def _standard_label(s: dict) -> str:
    return f"{s['framework']} {s['number']}" + (f"({s['name']})" if s.get("name") else "")


def _mark(i: int) -> str:
    return CIRCLED[i] if i < len(CIRCLED) else f"({i + 1})"


def _clip(text: str) -> str:
    return text if len(text) <= CELL_LIMIT else text[:CELL_LIMIT] + "\n…(이하 생략, '분류별 지적사항' 시트 참조)"


def _write_table(ws, start_row: int, headers: list[tuple[str, int]], rows: list[list], input_cols: set[int] = frozenset()):
    for col, (header, width) in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = FILL_INPUT if col in input_cols else FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[start_row].height = 30
    for r, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def _priority(findings: list[dict], this_year: int) -> tuple[float, str]:
    """(정렬용 점수, 표시 문구). 지적 빈도와 최근 3년 건수로 만든 힌트."""
    if not findings:
        return 0.0, "관련 지적사항 없음"
    years = [f["year"] for f in findings if f["year"]]
    recent = sum(1 for y in years if y >= this_year - 2)
    score = len(findings) + 2 * recent
    latest = f"최근 {max(years)}년" if years else "연도 미상"
    return score, f"지적 {len(findings)}건 · {latest} · 최근 3년 {recent}건"


def run_report(cfg: dict, store: Store, log, company: str | None = None) -> Path | None:
    companies = [r["company"] for r in store.query("SELECT DISTINCT company FROM account_mappings ORDER BY company")]
    if not companies:
        log.error("계정 매칭 결과가 없습니다. 먼저 'python main.py map --accounts <파일>'을 실행하세요.")
        return None
    if company is None:
        if len(companies) > 1:
            log.error("회사가 여러 개입니다. --company로 지정하세요: %s", ", ".join(companies))
            return None
        company = companies[0]
    elif company not in companies:
        log.error("'%s'의 계정 매칭 결과가 없습니다. 매칭된 회사: %s", company, ", ".join(companies))
        return None

    rcfg = cfg["report"]
    category_order = [c["name"] for c in cfg["taxonomy"]["categories"]]

    # ── 지적사항 선별: 제외 처리된 것과 근거(발췌)가 확인되지 않은 것은 싣지 않는다 ──
    all_findings = [dict(r) for r in store.query("SELECT * FROM findings ORDER BY year DESC, finding_id")]
    usable, excluded = {}, []
    for f in all_findings:
        if f["review_status"] == "제외":
            excluded.append((f, "검토에서 제외 처리"))
        elif rcfg["only_confirmed"] and f["review_status"] != "확정":
            excluded.append((f, "확정되지 않음(only_confirmed 설정)"))
        elif not f["source_excerpt"]:
            excluded.append((f, "원문 발췌 없음 — 출처 없는 위험 항목은 출력하지 않음"))
        elif f["excerpt_status"] not in VERIFIED and not rcfg["include_unverified_excerpt"]:
            excluded.append((f, f"발췌가 원문과 일치하지 않음({f['excerpt_status']}) — 검토 엑셀에서 발췌 수정 필요"))
        else:
            usable[f["finding_id"]] = f

    by_category: dict[str, list[dict]] = {}
    for r in store.query("SELECT finding_id, category FROM finding_categories"):
        if r["finding_id"] in usable:
            by_category.setdefault(r["category"], []).append(usable[r["finding_id"]])
    for items in by_category.values():
        items.sort(key=lambda f: (-(f["year"] or 0), f["finding_id"]))

    accounts = [dict(r) for r in store.query(
        "SELECT * FROM account_mappings WHERE company=? ORDER BY account_name", (company,))]
    accounts_of: dict[str, list[dict]] = {}
    for a in accounts:
        if a["category"]:
            accounts_of.setdefault(a["category"], []).append(a)

    signals = load_signals(store, company)
    this_year = datetime.now().year
    top_n = rcfg["max_findings_per_category"]

    # ── 시트 1: 계정별 체크리스트 (표준 분류별 1행) ──
    summary_rows, detail_rows = [], []
    categories = [c for c in category_order if c in accounts_of or c in by_category]
    for cat in categories:
        items = by_category.get(cat, [])
        accs = accounts_of.get(cat, [])
        score, hint = _priority(items, this_year)
        shown = items[:top_n]
        more = f"\n\n외 {len(items) - top_n}건 — '분류별 지적사항' 시트 참조" if len(items) > top_n else ""
        unreviewed = lambda f: " (미검토)" if f["review_status"] == "미검토" else ""  # noqa: E731
        standards = list(dict.fromkeys(
            _standard_label(s) for f in items for s in json.loads(f["standards"] or "[]")))
        vouchers = sum(a["n_vouchers"] or 0 for a in accs)
        # 계정이 많은 분류는 전표가 많은 계정부터 일부만 보이고 전체는 '계정 매칭' 시트에 둔다
        top_accs = sorted(accs, key=lambda a: -(a["n_vouchers"] or 0))[:rcfg["max_accounts_per_category"]]
        acc_text = ", ".join(a["account_name"] for a in top_accs)
        if len(accs) > len(top_accs):
            acc_text += f"\n외 {len(accs) - len(top_accs)}개 — '계정 매칭' 시트 참조"
        row = [
            cat,
            acc_text if accs else "(해당 계정 없음)",
            len(items),
            _clip("\n\n".join(f"{_mark(i)} [{f['finding_id']}]{unreviewed(f)} {f['risk_summary'] or ''}"
                              for i, f in enumerate(shown)) + more),
            "\n".join(standards),
            _clip("\n\n".join(f"{_mark(i)} [{f['finding_id']}] {f['audit_hint'] or ''}"
                              for i, f in enumerate(shown) if f["audit_hint"]) + more),
            _clip("\n\n".join(f"{_mark(i)} [{f['finding_id']}] {f['source_file']} p.{f['source_page']}\n“{f['source_excerpt']}”"
                              for i, f in enumerate(shown)) + more),
            hint,
            vouchers if accs else None,
        ]
        if signals:
            row.append(signals.get(cat, ""))
        row += ["미검토" if items else "해당없음", ""]
        summary_rows.append((score, len(accs) > 0, row))

        for f in items:
            detail_rows.append([
                cat, f["finding_id"], f["title"], f["issuer"], f["year"], f["finding_target"], f["finding_type"],
                f["risk_summary"], "\n".join(_standard_label(s) for s in json.loads(f["standards"] or "[]")),
                f["audit_hint"], f["source_excerpt"], f["excerpt_status"], f["source_file"], f["source_page"],
                f["review_status"],
            ])
    summary_rows.sort(key=lambda t: (-t[0], not t[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "계정별 체크리스트"
    ws.cell(row=1, column=1, value=f"감리 지적사항 대사 체크리스트 — {company}").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"작성일 {datetime.now():%Y-%m-%d} · 반영된 지적사항 {len(usable)}건"
                                     f"(미검토 {sum(1 for f in usable.values() if f['review_status'] == '미검토')}건 포함)")
    for i, line in enumerate(NOTICE, start=3):
        ws.cell(row=i, column=1, value=line).font = Font(color="C00000", bold=(i == 4))
    headers = [("계정과목(표준 분류)", 20), ("회사 계정", 34), ("관련 지적사항 수", 9), ("주요 위험", 72),
               ("관련 기준서", 30), ("권장 확인사항 (AI 제안, 감사절차 힌트)", 72), ("근거 (파일 / 페이지 / 원문 발췌)", 72),
               ("우선순위 힌트", 24), ("분개 전표 수 (참고)", 10)]
    if signals:
        headers.append(("위험 발현 가능성 (분개장)", 30))
    headers += [("감사인 검토 상태", 12), ("감사인 의견", 40)]
    n_cols = len(headers)
    _write_table(ws, 7, headers, [r for _, _, r in summary_rows], input_cols={n_cols - 1, n_cols})
    for r in range(8, 8 + len(summary_rows)):
        if ws.cell(row=r, column=3).value == 0:
            for c in range(1, n_cols + 1):
                ws.cell(row=r, column=c).fill = FILL_NONE
    if summary_rows:
        dv = DataValidation(type="list", formula1='"' + ",".join(AUDITOR_STATUSES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = get_column_letter(n_cols - 1)
        dv.add(f"{letter}8:{letter}{7 + len(summary_rows)}")
    ws.freeze_panes = "B8"

    # ── 시트 2: 분류별 지적사항 (분류 × 지적사항 1행, 전체 근거) ──
    _write_table(wb.create_sheet("분류별 지적사항"), 1, [
        ("표준 분류", 18), ("finding_id", 15), ("제목", 28), ("발행기관", 9), ("결정연도", 9), ("지적대상", 12),
        ("지적유형", 14), ("위험 요약", 60), ("관련 기준서", 30), ("권장 확인사항 (AI 제안)", 60),
        ("원문 발췌", 60), ("발췌 검증", 11), ("출처 파일", 32), ("페이지", 7), ("검토상태", 9)], detail_rows)

    # ── 시트 3: 지적사항 원장 (전체 구조화 레코드) ──
    categories_of: dict[str, list[str]] = {}
    for r in store.query("SELECT finding_id, category FROM finding_categories ORDER BY category"):
        categories_of.setdefault(r["finding_id"], []).append(r["category"])
    ledger = [[
        f["finding_id"], f["review_status"], f["title"], f["issuer"], f["year"], f["fiscal_period"], f["issue_area"],
        f["finding_target"], f["finding_type"],
        "\n".join(_standard_label(s) for s in json.loads(f["standards"] or "[]")),
        ", ".join(json.loads(f["related_accounts"] or "[]")), ", ".join(categories_of.get(f["finding_id"], [])),
        f["risk_summary"], f["source_excerpt"], f["excerpt_status"], f["audit_hint"], f["source_file"],
        f["source_page"], f["reviewer_note"], f["llm_model"],
    ] for f in all_findings]
    _write_table(wb.create_sheet("지적사항 원장"), 1, [
        ("finding_id", 15), ("검토상태", 9), ("제목", 28), ("발행기관", 9), ("결정연도", 9), ("회계결산일", 18),
        ("쟁점분야", 16), ("지적대상", 12), ("지적유형", 14), ("관련 기준서", 30), ("관련 계정", 24),
        ("표준 분류", 24), ("위험 요약", 60), ("원문 발췌", 60), ("발췌 검증", 11), ("권장 확인사항 (AI 제안)", 60),
        ("출처 파일", 32), ("페이지", 7), ("검토자 메모", 24), ("구조화 모델", 16)], ledger)

    # ── 시트 4: 계정 매칭 (회사 계정 전체와 분류 근거) ──
    _write_table(wb.create_sheet("계정 매칭"), 1, [
        ("표준 분류", 20), ("계정명", 34), ("계정코드", 12), ("매칭 방식", 10), ("신뢰도", 8), ("근거 키워드", 16),
        ("분개 전표 수 (참고)", 12)],
        [[a["category"], a["account_name"], a["account_code"], a["match_method"], a["match_confidence"],
          a["matched_keyword"], a["n_vouchers"]]
         for a in sorted((a for a in accounts if a["category"]),
                         key=lambda a: (category_order.index(a["category"]) if a["category"] in category_order else 999,
                                        a["account_name"]))])

    # ── 시트 5: 매칭 없음 계정 ──
    ws4 = wb.create_sheet("매칭 없음 계정")
    ws4.cell(row=1, column=1, value="표준 분류에 매칭되지 않은 계정입니다. 누락이 아니라 확인이 필요한 항목입니다. "
                                    "필요하면 config/account_taxonomy.yaml의 키워드나 overrides에 추가하세요.").font = Font(color="C00000")
    _write_table(ws4, 3, [("계정명", 34), ("계정코드", 12), ("분개 전표 수 (참고)", 12), ("감사인 의견", 40)],
                 [[a["account_name"], a["account_code"], a["n_vouchers"], ""] for a in accounts if not a["category"]],
                 input_cols={4})

    # ── 시트 6: 처리 로그 ──
    logs = [["체크리스트 제외", "WARNING", f["source_file"], f"{f['finding_id']} ({f['title']}): {reason}", ""]
            for f, reason in excluded]
    logs += [[r["stage"], r["level"], r["file_name"], r["message"], r["ts"]]
             for r in store.query("SELECT * FROM process_log ORDER BY (level='INFO'), id")]
    _write_table(wb.create_sheet("처리 로그"), 1,
                 [("단계", 14), ("수준", 10), ("파일", 40), ("내용", 100), ("시각", 19)], logs)

    safe = re.sub(r'[\\/:*?"<>|]', "_", company)
    path = cfg["paths"]["output"] / f"checklist_{safe}_{datetime.now():%Y%m%d}.xlsx"
    try:
        wb.save(path)
    except PermissionError:  # 같은 이름의 파일이 엑셀에서 열려 있는 경우
        path = path.with_name(f"{path.stem}_{datetime.now():%H%M%S}.xlsx")
        wb.save(path)
    log.info("체크리스트 출력: %s | 분류 %d행, 지적사항 %d건 반영, %d건 제외, 매칭 없음 계정 %d개",
             path, len(summary_rows), len(usable), len(excluded), sum(1 for a in accounts if not a["category"]))
    store.log("report", "INFO", path.name, f"{company} 체크리스트 출력 (지적사항 {len(usable)}건 반영, {len(excluded)}건 제외)")
    return path
