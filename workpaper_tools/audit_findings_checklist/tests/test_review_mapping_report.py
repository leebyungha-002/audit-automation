"""검토 엑셀 왕복, 계정 대사, 체크리스트 출력 테스트."""
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from afc.mapping import Taxonomy, _classify, load_account_list, run_map
from afc.report import run_report
from afc.review import COLUMNS, SHEET, export_review, import_review
from conftest import FakeLLM, add_finding, add_segment, add_source

COL = {header: i for i, (header, *_) in enumerate(COLUMNS, start=1)}


def edit_review(path, edits: dict[str, dict[str, object]]):
    wb = load_workbook(path)
    ws = wb[SHEET]
    rows = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
    for finding_id, changes in edits.items():
        for header, value in changes.items():
            ws.cell(row=rows[finding_id], column=COL[header], value=value)
    wb.save(path)


@pytest.fixture
def reviewed_env(env):
    cfg, store, log = env
    segment_id = add_segment(store, add_source(store))
    add_finding(store, "FSS-2106-03", segment_id)
    return cfg, store, log


def test_review_roundtrip_without_edits_changes_nothing(reviewed_env):
    cfg, store, log = reviewed_env
    path = export_review(cfg, store, log)
    import_review(cfg, store, log, path)
    f = store.query("SELECT * FROM findings")[0]
    assert f["reviewed_at"] is None and f["review_status"] == "미검토"


def test_review_edits_are_applied_and_ai_original_is_kept(reviewed_env):
    cfg, store, log = reviewed_env
    path = export_review(cfg, store, log)
    edit_review(path, {"FSS-2106-03": {
        "검토상태": "확정", "검토자 메모": "확인함", "관련 계정": "매출, 매출채권, 대여금",
        "관련 기준서": "K-IFRS 제1018호(수익)\n회계감사기준 500(감사증거)"}})
    import_review(cfg, store, log, path)

    f = store.query("SELECT * FROM findings")[0]
    assert f["review_status"] == "확정" and f["reviewer_note"] == "확인함" and f["reviewed_at"]
    assert json.loads(f["related_accounts"]) == ["매출", "매출채권", "대여금"]
    assert [s["check"] for s in json.loads(f["standards_check"])] == ["원문확인", "원문 미확인"]
    assert json.loads(f["llm_original_json"]) == {"original": True}


def test_review_rechecks_edited_excerpt_and_rejects_bad_values(reviewed_env):
    cfg, store, log = reviewed_env
    path = export_review(cfg, store, log)
    edit_review(path, {"FSS-2106-03": {"원문 발췌": "원문에 없는 문장", "검토상태": "보류", "관련 기준서": "형식 오류"}})
    import_review(cfg, store, log, path)

    f = store.query("SELECT * FROM findings")[0]
    assert f["excerpt_status"] == "불일치"
    assert f["review_status"] == "미검토"                                   # 허용값이 아니면 무시
    assert json.loads(f["standards"])[0]["number"] == "제1018호"            # 형식 오류면 기존 값 유지


@pytest.mark.parametrize("name, category", [
    ("외상매출금", "매출채권"),
    ("외상매출금대손충당금", "매출채권"),
    ("상품매출원가", "매출원가및제조원가"),
    ("반제품 매출", "매출"),
    ("제품매출-국내", "매출"),
    ("제품", "재고자산"),
    ("퇴직급여충당부채", "종업원급여"),
    ("이자비용(리스)", "리스"),
    ("사용권자산감가상각누계액", "리스"),
    ("투자부동산_건물감가상각누계액", "투자부동산"),
    ("건물 감가상각누계액", "유형자산"),
    ("미지급금-급여", "미지급금및미지급비용"),
    ("전환사채(유동)", "전환사채등복합금융상품"),
])
def test_dictionary_matching(env, name, category):
    assert Taxonomy(env[0]["taxonomy"]).match(name)[0] == category


def test_dictionary_returns_none_and_low_confidence_cases(env):
    tax = Taxonomy(env[0]["taxonomy"])
    assert tax.match("본지점(자산비용)") is None
    assert tax.match("보통예금")[1] == 1.0
    assert tax.match("건강보험_예수금")[1] < env[0]["mapping"]["llm_threshold"]  # 괄호·밑줄 뒤에서만 발견


def make_account_file(path):
    rows = [["회사명: 테스트상사", None, None, None, None, None],
            ["분석기간: 2025-01-01 ~ 2025-12-31", None, None, None, None, None],
            ["계정명", "차변합계", "차변건수", "대변합계", "대변건수", "전표개수"],
            ["보통예금", 123456789, 10, 987654321, 12, 20],
            ["[10800]외상매출금", 555, 3, 444, 2, 5],
            ["본지점(자산비용)", 1, 1, 1, 1, 1],
            ["0", 1, 1, 1, 1, 1]]
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name="05_계정명 리스트_통계", header=False, index=False)


def test_account_list_loading_reads_names_and_counts_but_not_amounts(tmp_path):
    path = tmp_path / "분석결과_test.xlsx"
    make_account_file(path)
    company, df = load_account_list(path)

    assert company == "테스트상사"
    assert list(df["account_name"]) == ["보통예금", "외상매출금", "본지점(자산비용)", "0"]
    assert df.loc[1, "account_code"] == "10800" and df.loc[0, "n_vouchers"] == 20
    assert not any("합계" in c for c in df.columns)


def test_only_account_names_are_sent_to_llm(env, tmp_path):
    cfg, store, log = env
    tax = Taxonomy(cfg["taxonomy"])
    fake = FakeLLM([{"results": [{"no": 1, "category": None, "confidence": "low"}]}])
    result = _classify(["보통예금", "본지점(자산비용)", "0"], tax, fake, cfg["mapping"]["llm_threshold"], log)

    sent = json.dumps(fake.calls[0]["content"], ensure_ascii=False)
    assert "본지점(자산비용)" in sent and "보통예금" not in sent   # 사전으로 확정된 계정은 보내지 않는다
    assert "1. " in sent and "0" not in sent.replace("1. ", "")    # 글자 없는 계정명도 보내지 않는다
    assert result["보통예금"]["method"] == "사전" and result["본지점(자산비용)"]["category"] is None


def test_map_stores_mappings_and_links_findings(reviewed_env, tmp_path):
    cfg, store, log = reviewed_env
    path = tmp_path / "분석결과_test.xlsx"
    make_account_file(path)
    run_map(cfg, store, log, path, use_llm=False)

    rows = {r["account_name"]: r for r in store.query("SELECT * FROM account_mappings")}
    assert rows["외상매출금"]["category"] == "매출채권" and rows["외상매출금"]["match_method"] == "사전"
    assert rows["본지점(자산비용)"]["match_method"] == "매칭 없음"
    assert {r["category"] for r in store.query("SELECT category FROM finding_categories")} == {"매출", "매출채권"}


def test_report_shows_only_findings_with_verified_evidence(env, tmp_path):
    cfg, store, log = env
    h = add_source(store)
    add_finding(store, "FSS-0001-01", add_segment(store, h, 1))
    add_finding(store, "FSS-0001-02", add_segment(store, h, 2), excerpt_status="불일치", title="발췌 불일치 사례")
    add_finding(store, "FSS-0001-03", add_segment(store, h, 3), review_status="제외", title="제외한 사례")
    add_finding(store, "FSS-0001-04", add_segment(store, h, 4), source_excerpt=None, excerpt_status="발췌없음")
    path = tmp_path / "분석결과_test.xlsx"
    make_account_file(path)
    run_map(cfg, store, log, path, use_llm=False)

    out = run_report(cfg, store, log, company="테스트상사")
    wb = load_workbook(out)
    ws = wb["계정별 체크리스트"]
    top = "\n".join(str(ws.cell(row=r, column=1).value) for r in range(1, 6))
    assert "최종 위험평가는 감사인의 판단" in top

    body = {ws.cell(row=r, column=1).value: [c.value for c in ws[r]] for r in range(8, ws.max_row + 1)}
    receivables = body["매출채권"]
    assert receivables[1] == "외상매출금" and receivables[2] == 1
    assert "[FSS-0001-01]" in receivables[3] and "FSS2106_03.pdf p.1" in receivables[6]
    everything = "\n".join(str(v) for row in body.values() for v in row)
    assert all(fid not in everything for fid in ("FSS-0001-02", "FSS-0001-03", "FSS-0001-04"))
    assert body["현금및예금"][2] == 0                                    # 지적사항 없는 분류도 행은 나온다

    unmatched = [r[0].value for r in wb["매칭 없음 계정"].iter_rows(min_row=4)]
    assert unmatched == ["0", "본지점(자산비용)"]
    excluded = [r[3].value for r in wb["처리 로그"].iter_rows(min_row=2) if r[0].value == "체크리스트 제외"]
    assert len(excluded) == 3
    assert len(list(wb["지적사항 원장"].iter_rows(min_row=2))) == 4      # 원장에는 전체가 남는다
    assert ws.data_validations.dataValidation                           # 감사인 검토 상태 드롭다운
